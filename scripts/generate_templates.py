"""Regenerate committed blank Office templates; not required to build the app.

Requires python-docx, openpyxl and python-pptx. The generated files contain no
user data and are copied directly by the app; Office need not be installed.
"""
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

destination = Path(__file__).resolve().parents[1] / "Resources" / "Templates"
destination.mkdir(parents=True, exist_ok=True)
document = Document()
document.core_properties.author = ""
document.core_properties.last_modified_by = ""
document.save(destination / "Blank.docx")
workbook = Workbook()
workbook.active.title = "Sheet1"
workbook.properties.creator = ""
workbook.properties.lastModifiedBy = ""
workbook.save(destination / "Blank.xlsx")
presentation = Presentation()
presentation.slides.add_slide(presentation.slide_layouts[6])
presentation.core_properties.author = ""
presentation.core_properties.last_modified_by = ""
presentation.save(destination / "Blank.pptx")
print("Generated three blank Office templates.")
